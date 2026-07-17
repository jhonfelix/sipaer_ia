import csv
import io
import logging
from datetime import datetime, timezone
from uuid import uuid4

import docx
import openpyxl
import pdfplumber
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from sqlalchemy import select

logger = logging.getLogger(__name__)

from app.database import AsyncSession, get_db
from app.middleware.auth import get_current_user
from app.models.conversation import Conversation
from app.models.project import Project
from app.models.user import User
from app.schemas.chat import (
    ChatRequest,
    ChatResponse,
    ConversationResponse,
    ConversationSessionResponse,
)
from app.agents.chat_agent import build_agent
from app.agents.tools.knowledge_tools import (
    collected_sources,
    get_tools,
    new_source_collector,
)
from app.services.rag_pipeline import (
    SYSTEM_PROMPT_GENERAL,
    SYSTEM_PROMPTS,
    rag_pipeline,
)

router = APIRouter(prefix="/chat", tags=["chat"])

# Limite de contexto anexado injetado nas instruções do agente (evita estourar a
# janela). O texto extraído já vem truncado pelo /chat/extract.
_AGENT_CTX_LIMIT = 120_000

# Instrução extra que orienta o agente a usar as tools de busca e citar fontes.
_AGENT_TOOL_GUIDANCE = (
    "\n\nVocê dispõe de ferramentas de busca na base de conhecimento SIPAER. "
    "Use-as para fundamentar respostas factuais e cite as fontes encontradas. "
    "Se a base não retornar nada relevante, responda com seu conhecimento e "
    "sinalize que não havia fonte na base."
)

_EXTRACTABLE = {
    "application/pdf",
    "application/msword",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/vnd.ms-excel",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "text/plain",
    "text/csv",
}
_EXTRACT_MAX_BYTES = 300 * 1024 * 1024  # 300 MB
_EXTRACT_MAX_CHARS = 200_000


@router.post("/extract")
async def extract_text(
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
) -> dict:
    if file.content_type not in _EXTRACTABLE:
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail=f"Tipo não suportado para extração: {file.content_type}",
        )

    content = await file.read()
    if len(content) > _EXTRACT_MAX_BYTES:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail="Arquivo excede 300 MB",
        )

    mime = file.content_type or ""
    text = ""

    if mime == "application/pdf":
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            text = "\n\n".join(p.extract_text() or "" for p in pdf.pages)

    elif "wordprocessingml" in mime or mime == "application/msword":
        doc = docx.Document(io.BytesIO(content))
        text = "\n".join(p.text for p in doc.paragraphs)

    elif "spreadsheetml" in mime or mime == "application/vnd.ms-excel":
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        rows: list[str] = []
        for ws in wb.worksheets:
            rows.append(f"[Planilha: {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                rows.append("\t".join("" if v is None else str(v) for v in row))
        text = "\n".join(rows)

    elif mime == "text/csv":
        decoded = content.decode("utf-8", errors="replace")
        reader = csv.reader(io.StringIO(decoded))
        text = "\n".join("\t".join(cell for cell in row) for row in reader)

    else:  # text/plain
        text = content.decode("utf-8", errors="replace")

    truncated = len(text) > _EXTRACT_MAX_CHARS
    return {
        "text": text[:_EXTRACT_MAX_CHARS],
        "filename": file.filename or "arquivo",
        "truncated": truncated,
    }


@router.post("", response_model=ChatResponse)
async def chat(
    body: ChatRequest,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    session_id = body.session_id or str(uuid4())
    now = datetime.now(timezone.utc)

    # Projeto (opcional): valida posse e carrega as instruções para o contexto.
    project_instructions: str | None = None
    if body.project_id is not None:
        result = await db.execute(
            select(Project).where(
                Project.id == body.project_id,
                Project.created_by == current_user.id,
            )
        )
        project = result.scalar_one_or_none()
        if not project:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Projeto não encontrado")
        project_instructions = project.instructions

    user_msg = Conversation(
        user_id=current_user.id,
        report_id=body.report_id,
        project_id=body.project_id,
        session_id=session_id,
        category=body.chat_type,
        role="user",
        content=body.message,
        sources=[],
    )
    db.add(user_msg)
    await db.flush()

    try:
        content, sources = await rag_pipeline.process(
            body.message,
            body.context,
            body.model,
            body.chat_type,
            project_id=body.project_id,
            project_instructions=project_instructions,
        )
    except Exception:
        logger.exception("RAG pipeline falhou")
        content = "Serviço de IA temporariamente indisponível. Tente novamente em instantes."
        sources = []

    ai_msg = Conversation(
        user_id=current_user.id,
        report_id=body.report_id,
        project_id=body.project_id,
        session_id=session_id,
        category=body.chat_type,
        role="assistant",
        content=content,
        sources=sources,
    )
    db.add(ai_msg)
    await db.commit()
    await db.refresh(ai_msg)

    return ChatResponse(
        id=str(ai_msg.id),
        content=content,
        sources=sources,
        session_id=session_id,
        created_at=now,
    )


async def _run_chat_agent(
    body: ChatRequest,
    session_id: str,
    user: User,
    project_instructions: str | None,
) -> tuple[str, list[dict]]:
    """Executa o agente Agno (tool calling + memória) e devolve (conteúdo, fontes)."""
    instructions = SYSTEM_PROMPTS.get(body.chat_type, SYSTEM_PROMPT_GENERAL)
    instructions += _AGENT_TOOL_GUIDANCE
    if project_instructions and project_instructions.strip():
        instructions += f"\n\nInstruções do projeto:\n{project_instructions.strip()}"
    if body.context:
        instructions += f"\n\nContexto anexado pelo usuário:\n{body.context[:_AGENT_CTX_LIMIT]}"

    # Coletor de fontes por-requisição (as tools fazem append via contextvars).
    new_source_collector()
    agent = build_agent(
        model_id=body.model,
        instructions=instructions,
        tools=get_tools(body.chat_type),
        session_id=session_id,
        user_id=str(user.id),
        enable_memory=True,
    )
    result = await agent.arun(body.message)
    content = getattr(result, "content", None) or ""
    return content, collected_sources()


@router.post("/agent", response_model=ChatResponse)
async def chat_agent(
    body: ChatRequest,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Chat agêntico (Agno) — coexiste com POST /chat clássico.

    O agente decide o que/onde buscar (tool calling) e tem memória de sessão
    (MySQL + Agno). `images` e chats de projeto ainda usam o rag_pipeline clássico.
    """
    session_id = body.session_id or str(uuid4())
    now = datetime.now(timezone.utc)

    project_instructions: str | None = None
    if body.project_id is not None:
        result = await db.execute(
            select(Project).where(
                Project.id == body.project_id,
                Project.created_by == current_user.id,
            )
        )
        project = result.scalar_one_or_none()
        if not project:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Projeto não encontrado")
        project_instructions = project.instructions

    user_msg = Conversation(
        user_id=current_user.id,
        report_id=body.report_id,
        project_id=body.project_id,
        session_id=session_id,
        category=body.chat_type,
        role="user",
        content=body.message,
        sources=[],
    )
    db.add(user_msg)
    await db.flush()

    # O agente cobre os chats de base de conhecimento. `images` (visão) e chats de
    # projeto (busca escopada por project_id) permanecem no pipeline clássico.
    use_agent = body.chat_type != "images" and body.project_id is None

    try:
        if use_agent:
            content, sources = await _run_chat_agent(
                body, session_id, current_user, project_instructions
            )
        else:
            content, sources = await rag_pipeline.process(
                body.message,
                body.context,
                body.model,
                body.chat_type,
                project_id=body.project_id,
                project_instructions=project_instructions,
            )
    except Exception:
        logger.exception("Chat agent falhou")
        content = "Serviço de IA temporariamente indisponível. Tente novamente em instantes."
        sources = []

    ai_msg = Conversation(
        user_id=current_user.id,
        report_id=body.report_id,
        project_id=body.project_id,
        session_id=session_id,
        category=body.chat_type,
        role="assistant",
        content=content,
        sources=sources,
    )
    db.add(ai_msg)
    await db.commit()
    await db.refresh(ai_msg)

    return ChatResponse(
        id=str(ai_msg.id),
        content=content,
        sources=sources,
        session_id=session_id,
        created_at=now,
    )


@router.get("/sessions", response_model=list[ConversationSessionResponse])
async def sessions(
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(Conversation)
        .where(
            Conversation.user_id == current_user.id,
            Conversation.session_id.isnot(None),
        )
        .order_by(Conversation.session_id, Conversation.created_at)
    )
    rows = result.scalars().all()

    sessions_dict: dict[str, list[Conversation]] = {}
    for row in rows:
        sessions_dict.setdefault(row.session_id, []).append(row)  # type: ignore[index]

    out: list[ConversationSessionResponse] = []
    for sid, msgs in sessions_dict.items():
        first_user = next((m for m in msgs if m.role == "user"), None)
        last_msg = msgs[-1]
        raw_title = first_user.content if first_user else "Conversa"
        title = raw_title[:50] + ("…" if len(raw_title) > 50 else "")
        preview = last_msg.content[:60] + ("…" if len(last_msg.content) > 60 else "")
        category = first_user.category or "general" if first_user else "general"
        out.append(
            ConversationSessionResponse(
                session_id=sid,
                title=title,
                preview=preview,
                message_count=len(msgs),
                updated_at=last_msg.created_at,
                category=category,
            )
        )

    out.sort(key=lambda s: s.updated_at, reverse=True)
    return out


@router.get("/history", response_model=list[ConversationResponse])
async def history(
    session_id: str | None = None,
    report_id: int | None = None,
    project_id: int | None = None,
    limit: int = 100,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    query = (
        select(Conversation)
        .where(Conversation.user_id == current_user.id)
        .order_by(Conversation.created_at)
        .limit(limit)
    )
    if session_id is not None:
        query = query.where(Conversation.session_id == session_id)
    if report_id is not None:
        query = query.where(Conversation.report_id == report_id)
    if project_id is not None:
        query = query.where(Conversation.project_id == project_id)

    result = await db.execute(query)
    return list(result.scalars().all())
